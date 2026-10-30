"""Regression tests for the practicum workspace, importer, learning, PDF."""

import copy
import io
import json
import unittest

from openpyxl import load_workbook

from utils.education import LESSONS, worked_example
from utils.importer import make_excel_template, read_excel_observations
from utils.module_report import export_module_pdf
from utils.workspace import create_project_file, read_project_file, cloud_save, cloud_load


def sample_state():
    """Small realistic source data with all backup fields, no Streamlit required."""
    from utils.exporter import SESSION_BACKUP_FIELDS
    values = {key: {} for _, key in SESSION_BACKUP_FIELDS}
    values["identitas"] = {"universitas": "Universitas", "kelas": "J253B", "kelompok": "020",
                            "tanggal": "30-10-2026", "dosen": "Dosen", "op1_usia": 21,
                            "op1_jk": "Laki-laki"}
    values["modul_01"] = {"op1_akt1_dni": 80.0, "op1_akt1_dnk": 110.0}
    values["modul_02"] = {"records": [{"rula_params": {
        "ua": 3, "la": 2, "wr": 2, "wt": 1, "nk": 3, "tr": 3, "lg": 1,
        "ma": 1, "fa": 1, "mb": 1, "fb": 1},
        "owas_params": {"back": 2, "arms": 1, "legs": 2, "load": 1},
        "reba_params": {"neck": 2, "trunk": 3, "legs": 1, "load": 1,
                        "ua": 3, "la": 2, "wr": 2, "coup": 1, "act": 1}}]}
    values["modul_03"] = {"df_siklus": [{"Kode": "EL-1", "Elemen Kerja": "Menulis",
                             **{f"S{i}": 20.0 + i for i in range(1, 11)}}],
                           "westinghouse": {"skill": 0.06, "effort": 0.05, "condition": 0.02, "consistency": 0.01},
                           "allowance_pribadi": 5.0, "allowance_kelelahan": 7.0,
                           "allowance_hambatan": 3.0, "k_level": 2.0, "s_level": 0.05}
    values["modul_04"] = {"df_32": [{"Kode": "D01", "Dimensi Tubuh": "Tinggi tubuh",
                              "Aplikasi Desain": "P95", **{f"P{i}": 160.0 + i for i in range(1, 11)}}]}
    values["modul_05"] = {"df_3dim": [{"Dimensi": "Tinggi popliteal", "Persentil Acuan": "P5",
                               "Kelonggaran (cm)": 0.0, **{f"S{i}": 40.0 + i for i in range(1, 11)}}]}
    values["modul_06"] = {"dingin_suhu": [21.0, 21.5, 22.0], "panas_suhu": [30.0, 31.0, 32.0]}
    values["modul_07"] = {"normal_db": [60.0, 62.0, 61.0], "bising_db": [87.0, 88.0, 89.0]}
    values["modul_08"] = {"normal_lux": [300.0, 320.0, 310.0], "gelap_lux": [80.0, 90.0, 100.0]}
    values["modul_09"] = {"h_cm": 32.0, "v_cm": 60.0, "d_cm": 45.0, "a_deg": 15.0,
                           "f_lift": 2.0, "durasi_angkat": "≤ 1 jam", "coupling": "Sedang (Fair)",
                           "berat_beban": 16.0}
    values["module_done"] = {"modul-01": True}
    values["analisis_modul"] = {"Modul 03": "Perlu pengukuran tambahan."}
    return values


class RoadmapFeatureTests(unittest.TestCase):
    def test_project_file_roundtrip_including_progress(self):
        state = sample_state()
        filename = create_project_file("  Kelompok   020  ", state)
        name, backup = read_project_file(filename.encode("utf-8"))
        self.assertEqual(name, "Kelompok 020")
        self.assertTrue(backup["module_done"]["modul-01"])
        self.assertEqual(backup["modul_03_waktu_kerja"]["df_siklus"], state["modul_03"]["df_siklus"])

    def test_bad_project_rejected_without_state_mutation(self):
        state = sample_state()
        old = copy.deepcopy(state)
        with self.assertRaises(ValueError):
            read_project_file(b'{"format":"ie-toolkit-project-v1","name":"salah","backup":{"identitas":{}}}')
        with self.assertRaises(ValueError):
            read_project_file(b"not json")
        with self.assertRaises(ValueError):
            read_project_file(b"x" * (5 * 1024 * 1024 + 1))
        with self.assertRaises(ValueError):
            create_project_file(" ", state)
        self.assertEqual(state, old)

    def test_cloud_writes_without_owner_are_blocked(self):
        with self.assertRaises(ValueError):
            cloud_save("postgres://example", "", "A", {})
        with self.assertRaises(ValueError):
            cloud_load("postgres://example", "", "A")

    def test_excel_import_03_roundtrip(self):
        state = sample_state()
        blob = make_excel_template("03", state)
        updated = read_excel_observations("03", blob, state["modul_03"])
        self.assertEqual(updated["df_siklus"], state["modul_03"]["df_siklus"])

    def test_excel_import_05_keeps_numeric_allowance(self):
        state = sample_state()
        blob = make_excel_template("05", state)
        updated = read_excel_observations("05", blob, state["modul_05"])
        self.assertIsInstance(updated["df_3dim"][0]["Kelonggaran (cm)"], float)
        self.assertEqual(updated["df_3dim"][0]["Kelonggaran (cm)"], 0.0)

    def test_excel_import_all_environment_modules(self):
        state = sample_state()
        for mod in ("06", "07", "08"):
            with self.subTest(mod=mod):
                blob = make_excel_template(mod, state)
                updated = read_excel_observations(mod, blob, state[f"modul_{mod}"])
                self.assertEqual(updated, state[f"modul_{mod}"])

    def test_excel_rejects_formula_bad_column_and_out_of_range_atomically(self):
        state = sample_state()
        original = copy.deepcopy(state["modul_06"])
        for replacement in ("=2+3", -3, float("inf"), 46):
            with self.subTest(value=replacement):
                wb = load_workbook(io.BytesIO(make_excel_template("06", state)))
                wb["Data"]["B2"] = replacement
                output = io.BytesIO()
                wb.save(output)
                with self.assertRaises(ValueError):
                    read_excel_observations("06", output.getvalue(), state["modul_06"])
                self.assertEqual(state["modul_06"], original)
        wb = load_workbook(io.BytesIO(make_excel_template("06", state)))
        wb["Data"]["B1"] = "Bukan Kolom"
        output = io.BytesIO()
        wb.save(output)
        with self.assertRaises(ValueError):
            read_excel_observations("06", output.getvalue(), state["modul_06"])

    def test_every_module_has_learning_material_and_trace(self):
        state = sample_state()
        self.assertEqual(len(LESSONS), 9)
        for module in LESSONS:
            with self.subTest(module=module):
                self.assertTrue(LESSONS[module][2])
                self.assertTrue(worked_example(module, state))
        self.assertIn("25.21", worked_example("01", state)[1])
        self.assertIn("15", worked_example("03", state)[3])

    def test_module_pdf_valid_for_all_nine_modules(self):
        state = sample_state()
        result = {"m1_rows": [], "m2_rows": [], "m3_rows": [], "m4_rows": [],
                  "m5_rows": [], "env_rows": [["Modul 06"], ["Modul 07"], ["Modul 08"]],
                  "m9_summary": {"rwl": 17, "li": 0.94, "f_kompresi_n": 2800}}
        for module in LESSONS:
            with self.subTest(module=module):
                pdf = export_module_pdf(module, state, result)
                self.assertTrue(pdf.startswith(b"%PDF"))
                self.assertGreater(len(pdf), 900)


if __name__ == "__main__":
    unittest.main()
