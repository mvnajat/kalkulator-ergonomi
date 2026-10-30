import io
import json
from collections.abc import MutableMapping
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


SESSION_BACKUP_FIELDS = (
    ("identitas", "identitas"),
    ("modul_01_fisiologi", "modul_01"),
    ("modul_02_postur", "modul_02"),
    ("modul_03_waktu_kerja", "modul_03"),
    ("modul_04_antropometri_32", "modul_04"),
    ("modul_05_antropometri_lanjutan", "modul_05"),
    ("modul_06_temperatur", "modul_06"),
    ("modul_07_kebisingan", "modul_07"),
    ("modul_08_pencahayaan", "modul_08"),
    ("modul_09_mmh", "modul_09"),
    ("modul_09_simulasi", "modul_09_simulasi"),
    ("analisis_modul", "analisis_modul"),
    ("module_done", "module_done"),
)


def export_session_to_json(state_dict: dict) -> str:
    """Mengekspor seluruh data praktikum dari session_state ke format JSON terstruktur."""
    export_data = {json_key: state_dict.get(state_key, {}) for json_key, state_key in SESSION_BACKUP_FIELDS}
    return json.dumps(export_data, indent=2, ensure_ascii=False)


def restore_session_from_json(state: MutableMapping, backup: dict) -> None:
    """Pulihkan bidang yang tersedia, termasuk backup lama tanpa Modul 09."""
    if not isinstance(backup, dict):
        raise ValueError("Format cadangan harus berupa objek JSON.")

    restored = {}
    for json_key, state_key in SESSION_BACKUP_FIELDS:
        if json_key in backup:
            if not isinstance(backup[json_key], dict):
                raise ValueError(f"Data '{json_key}' harus berupa objek JSON.")
            restored[state_key] = backup[json_key]

    for state_key, value in restored.items():
        state[state_key] = value


def export_session_to_excel(state_dict: dict, calc_results: dict) -> bytes:
    """Membuat file laporan Excel (.xlsx) lengkap dengan styling rapi untuk seluruh modul."""
    wb = openpyxl.Workbook()
    # Hapus sheet default
    wb.remove(wb.active)

    # Palet Styling
    header_fill = PatternFill(start_color="0F2E1E", end_color="0F2E1E", fill_type="solid")
    sub_fill = PatternFill(start_color="19452F", end_color="19452F", fill_type="solid")
    zebra_fill = PatternFill(start_color="F2F8F5", end_color="F2F8F5", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=14, bold=True, color="0F2E1E")
    bold_font = Font(name="Calibri", size=10, bold=True)
    reg_font = Font(name="Calibri", size=10)
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    def style_worksheet(ws, title_text: str):
        ws.views.sheetView[0].showGridLines = True
        ws["A1"] = title_text
        ws["A1"].font = title_font
        ws["A2"] = "UNIVERSITAS PELITA BANGSA - PROGRAM STUDI TEKNIK INDUSTRI"
        ws["A2"].font = bold_font

    def write_table(ws, start_row: int, headers: list, rows: list):
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=start_row, column=col_idx, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border
        for r_idx, row_data in enumerate(rows, start=start_row + 1):
            for col_idx, val in enumerate(row_data, start=1):
                cell = ws.cell(row=r_idx, column=col_idx, value=val)
                # Teks bebas dari pengguna tidak boleh berubah menjadi rumus Excel.
                if isinstance(val, str) and val.startswith(("=", "+", "-", "@")):
                    cell.data_type = "s"
                cell.font = reg_font
                cell.border = thin_border
                if isinstance(val, (int, float)) and not isinstance(val, bool):
                    cell.number_format = "0.00"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if (r_idx - start_row) % 2 == 0:
                    cell.fill = zebra_fill
        # Auto-fit lebar kolom
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                max_len = max(max_len, len(val_str))
            ws.column_dimensions[col_letter].width = min(max(max_len + 4, 14), 45)

    # 1. Sheet Identitas & Ringkasan
    ws_id = wb.create_sheet("00_Identitas")
    style_worksheet(ws_id, "IDENTITAS KELOMPOK PRAKTIKUM ERGONOMI")
    id_data = state_dict.get("identitas", {})
    rows_id = [
        ["Kelas", id_data.get("kelas", "J253B")],
        ["Kelompok", id_data.get("kelompok", "Kelompok 1")],
        ["Tanggal Praktikum", id_data.get("tanggal", "")],
        ["Asisten / Dosen Pengampu", id_data.get("dosen", "")],
        ["Operator 1 (Nama & Usia)", f"{id_data.get('op1_nama', '')} ({id_data.get('op1_usia', '')} thn)"],
        ["Operator 2 (Nama & Usia)", f"{id_data.get('op2_nama', '')} ({id_data.get('op2_usia', '')} thn)"],
        ["Daftar Anggota Kelompok", id_data.get("anggota", "")],
    ]
    write_table(ws_id, 4, ["Parameter Identitas", "Keterangan"], rows_id)

    # 2. Sheet Modul 01 Fisiologi
    ws_m1 = wb.create_sheet("01_Fisiologi")
    style_worksheet(ws_m1, "MODUL 01 - FISIOLOGI KERJA & BEBAN KARDIOVASKULAR (%CVL)")
    m1_rows = calc_results.get("m1_rows", [])
    write_table(ws_m1, 4, [
        "Operator", "Aktivitas", "Usia (th)", "DN Istirahat (bpm)", "DN Kerja (bpm)",
        "DN Maks (bpm)", "%CVL (%)", "Konsumsi Energi Bersih (kkal/mnt)", "Kategori Beban", "Rekomendasi"
    ], m1_rows)
    if calc_results.get("m1_recovery_rows"):
        ws_pulse = wb.create_sheet("01_Pemulihan_Nadi")
        style_worksheet(ws_pulse, "MODUL 01 - PEMULIHAN NADI P0-P5 / INDEKS ECPT-ECPM")
        write_table(ws_pulse, 4, [
            "Operator", "P0 (bpm)", "P1 (bpm)", "P2 (bpm)", "P3 (bpm)",
            "P4 (bpm)", "P5 (bpm)", "ECPT (bpm)", "ECPM (bpm)",
            "Pemulihan P1 ke P5 (%)", "Catatan Interpretasi",
        ], calc_results["m1_recovery_rows"])
        ws_pulse["A3"] = "ECPT/ECPM adalah indeks selisih nadi (bpm), bukan nilai energi kkal."

    # 3. Sheet Modul 02 Postur (RULA, OWAS, REBA)
    ws_m2 = wb.create_sheet("02_Postur_Kerja")
    style_worksheet(ws_m2, "MODUL 02 - ANALISIS POSTUR KERJA (RULA, OWAS, & REBA)")
    m2_rows = calc_results.get("m2_rows", [])
    write_table(ws_m2, 4, [
        "Operator", "Aktivitas", "Metode", "Skor / Kode Akhir", "Action Level", "Kategori Risiko", "Rekomendasi Perbaikan"
    ], m2_rows)

    # 4. Sheet Modul 03 Waktu Kerja
    ws_m3 = wb.create_sheet("03_Waktu_Kerja")
    style_worksheet(ws_m3, "MODUL 03 - PENGUKURAN WAKTU KERJA (TIME STUDY)")
    m3_rows = calc_results.get("m3_rows", [])
    write_table(ws_m3, 4, [
        "Kode", "Elemen Kerja", "Ws Rata-rata (dtk)", "SD (dtk)", "BKB", "BKA",
        "Keseragaman", "N'", "Kecukupan", "Faktor p", "Allowance (%)", "Wn (dtk)", "Wb (dtk)", "Output Standar (unit/jam)"
    ], m3_rows)
    if calc_results.get("m3_comparison_rows"):
        ws_rates = wb.create_sheet("03_Perbandingan_Metode")
        style_worksheet(ws_rates, "MODUL 03 - SCHUMMARD, WESTINGHOUSE, OBJEKTIF DAN AKHIR")
        ws_rates["A3"] = "Faktor p: Schummard=skor/60; Objektif=p1*(1+total/100); Akhir=manual."
        write_table(ws_rates, 4, [
            "Kode", "Elemen Kerja", "Metode", "Ws (dtk)", "Faktor p", "Schummard p1",
            "Objektif p2", "Total Kesulitan", "Allowance (%)", "Wn (dtk)", "Wb (dtk)",
            "Output per Elemen (unit/jam)", "N'", "Kecukupan", "Keseragaman",
        ], calc_results["m3_comparison_rows"])

    # 5. Sheet Modul 04 Antropometri 32 Dimensi
    ws_m4 = wb.create_sheet("04_Antropometri_32D")
    style_worksheet(ws_m4, "MODUL 04 - ANTROPOMETRI (32 DIMENSI TUBUH & PERSENTIL DESAIN)")
    m4_rows = calc_results.get("m4_rows", [])
    write_table(ws_m4, 4, [
        "Kode", "Nama Dimensi Tubuh", "Mean (cm)", "SD (cm)", "P5 (cm)", "P10 (cm)", "P50 (cm)", "P90 (cm)", "P95 (cm)", "Uji Normalitas", "Aplikasi Desain Stasiun Kerja"
    ], m4_rows)

    # 6. Sheet Modul 05 Antropometri Lanjutan
    ws_m5 = wb.create_sheet("05_Antropometri_Lanjutan")
    style_worksheet(ws_m5, "MODUL 05 - ANTROPOMETRI LANJUTAN (UJI KESERAGAMAN, KECUKUPAN & RANCANGAN)")
    m5_rows = calc_results.get("m5_rows", [])
    write_table(ws_m5, 4, [
        "Dimensi Kunci", "Mean (cm)", "SD (cm)", "BKB (cm)", "BKA (cm)",
        "Status Keseragaman", "N'", "Status Kecukupan", "Persentil Acuan", "Ukuran Rancangan (cm)"
    ], m5_rows)

    # 7. Sheet Modul 06, 07, 08 Lingkungan Fisik
    ws_env = wb.create_sheet("06_08_Lingkungan_Kerja")
    style_worksheet(ws_env, "MODUL 06, 07, & 08 - LINGKUNGAN FISIK KERJA (TEMPERATUR, KEBISINGAN, PENCAHAYAAN)")
    env_rows = calc_results.get("env_rows", [])
    write_table(ws_env, 4, [
        "Modul", "Parameter Uji", "Rata-rata Kondisi Normal/Dingin", "Rata-rata Kondisi Ekstrem/Bising/Gelap",
        "Selisih", "Acuan Perbandingan Praktikum", "p-value (Uji-t)", "Kesimpulan Statistik", "Analisis & Rekomendasi"
    ], env_rows)

    # 8. Sheet Modul 09, termasuk input lengkap dan hasil perhitungan.
    ws_m9 = wb.create_sheet("09_MMH")
    style_worksheet(ws_m9, "MODUL 09 - MANUAL MATERIAL HANDLING (NIOSH, SNOOK, BIOMEKANIKA)")
    m9 = state_dict.get("modul_09", {})
    summary = calc_results.get("m9_summary", {})
    labels_m9 = [
        ("h_cm", "Jarak horizontal H (cm)"), ("v_cm", "Tinggi awal V (cm)"),
        ("d_cm", "Jarak angkat D (cm)"), ("a_deg", "Sudut asimetri A (derajat)"),
        ("f_lift", "Frekuensi (kali/menit)"), ("durasi_angkat", "Durasi pengangkatan"),
        ("coupling", "Coupling"), ("berat_beban", "Berat beban (kg)"),
        ("snook_aktivitas", "Aktivitas Snook (simulasi)"),
        ("snook_jk", "Populasi pekerja (simulasi)"),
        ("snook_persentil", "Persentil (simulasi)"),
        ("bb_pekerja", "Berat badan pekerja (kg)"),
        ("sudut_punggung", "Sudut punggung (derajat)"),
    ]
    m9_rows = [[label, m9.get(key, "-")] for key, label in labels_m9]
    m9_rows += [
        ["NIOSH RWL (kg)", summary.get("rwl")],
        ["NIOSH Lifting Index", summary.get("li")],
        ["Kategori NIOSH", summary.get("kategori", "-")],
        ["Estimasi kompresi L5/S1 (N)", summary.get("f_kompresi_n")],
        ["Status model L5/S1", summary.get("status_l5s1", "-")],
        ["Keterbatasan", "Tabel Snook hanya simulasi; L5/S1 model sederhana; penilaian akhir memerlukan pengukuran ergonomi nyata."],
    ]
    write_table(ws_m9, 4, ["Parameter", "Nilai / Hasil"], m9_rows)

    sim_data = state_dict.get("modul_09_simulasi")
    sim_summary = calc_results.get("m9_simulasi_summary")
    if isinstance(sim_data, dict) and isinstance(sim_summary, dict):
        ws_sim = wb.create_sheet("09_Simulasi")
        style_worksheet(ws_sim, "MODUL 09 - PERBANDINGAN KONDISI ANGKAT DAN USULAN PERBAIKAN")
        sim_fields = labels_m9[:8]
        sim_rows = [[label, m9.get(key), sim_data.get(key)] for key, label in sim_fields]
        sim_rows.extend([
            ["Recommended Weight Limit (kg)", sim_summary["awal"]["rwl"], sim_summary["usulan"]["rwl"]],
            ["Lifting Index", sim_summary["awal"]["li"], sim_summary["usulan"]["li"]],
            ["Kategori", sim_summary["awal"]["kategori"], sim_summary["usulan"]["kategori"]],
            ["Perubahan LI (%)", None, sim_summary["perubahan_li_pct"]],
            ["Interpretasi", "Perbandingan edukatif, bukan kepastian bebas cedera", sim_summary["status"]],
        ])
        write_table(ws_sim, 4, ["Parameter", "Kondisi Awal", "Usulan"], sim_rows)

    # Sheet pengamatan mentah membuat angka ringkasan dapat diperiksa ulang.
    def write_observations(sheet_name, title, records, key_label, second_label, sample_prefix):
        ws = wb.create_sheet(sheet_name)
        style_worksheet(ws, title)
        sample_names = [f"{sample_prefix}{i}" for i in range(1, 11)]
        rows = []
        for rec in records:
            if not isinstance(rec, dict):
                continue
            if sheet_name == "05_Data_Mentah":
                metadata = f"Acuan {rec.get('Persentil Acuan', '-')}; kelonggaran {rec.get('Kelonggaran (cm)', 0)} cm"
            elif sheet_name == "04_Data_Mentah":
                metadata = rec.get("Aplikasi Desain", "")
            else:
                metadata = ""
            rows.append([rec.get(key_label, ""), rec.get(second_label, "")]
                        + [rec.get(col) for col in sample_names]
                        + [None, None, metadata])
        write_table(ws, 4, ["Kode / Dimensi", "Nama / Elemen"] + sample_names +
                    ["Rata-rata (rumus)", "SD Sampel (rumus)", "Keterangan"], rows)
        for row_number in range(5, 5 + len(rows)):
            sample_range = f"C{row_number}:L{row_number}"
            ws[f"M{row_number}"] = f'=IF(COUNT({sample_range})=0,"",AVERAGE({sample_range}))'
            ws[f"N{row_number}"] = f'=IF(COUNT({sample_range})<2,"",STDEV.S({sample_range}))'
            ws[f"M{row_number}"].number_format = "0.00"
            ws[f"N{row_number}"].number_format = "0.00"
        ws.freeze_panes = "C5"

    m3 = state_dict.get("modul_03", {})
    m4 = state_dict.get("modul_04", {})
    m5 = state_dict.get("modul_05", {})
    write_observations("03_Data_Mentah", "MODUL 03 - 10 OBSERVASI WAKTU SIKLUS (detik)",
                       m3.get("df_siklus", []), "Kode", "Elemen Kerja", "S")
    write_observations("04_Data_Mentah", "MODUL 04 - 10 SAMPEL ANTROPOMETRI (cm)",
                       m4.get("df_32", []), "Kode", "Dimensi Tubuh", "P")
    write_observations("05_Data_Mentah", "MODUL 05 - 10 SAMPEL DIMENSI RANCANGAN (cm)",
                       m5.get("df_3dim", []), "Dimensi", "Dimensi", "S")

    ws_env_raw = wb.create_sheet("06_08_Data_Mentah")
    style_worksheet(ws_env_raw, "MODUL 06-08 - HASIL PENGUKURAN (masing-masing 3 percobaan)")
    observations = []
    for module_key, fields in (
        ("modul_06", (("dingin_suhu", "Temperatur dingin (°C)"), ("panas_suhu", "Temperatur panas (°C)"),
                      ("dingin_waktu", "Waktu kondisi dingin (detik)"), ("panas_waktu", "Waktu kondisi panas (detik)"))),
        ("modul_07", (("normal_db", "Kebisingan normal (dBA)"), ("bising_db", "Kebisingan bising (dBA)"),
                      ("normal_waktu", "Waktu kondisi normal (detik)"), ("bising_waktu", "Waktu kondisi bising (detik)"))),
        ("modul_08", (("normal_lux", "Cahaya normal (lux)"), ("gelap_lux", "Cahaya redup (lux)"),
                      ("normal_waktu", "Waktu kondisi normal (detik)"), ("gelap_waktu", "Waktu kondisi redup (detik)"))),
    ):
        for key, label in fields:
            values = state_dict.get(module_key, {}).get(key, [])
            observations.append([module_key.replace("modul_", "Modul "), label] + list(values[:3]))
    write_table(ws_env_raw, 4, ["Modul", "Parameter", "Percobaan 1", "Percobaan 2", "Percobaan 3"], observations)

    # Data mentah ditempatkan di samping ringkasan yang relevan.
    order = ["00_Identitas", "01_Fisiologi"]
    if "01_Pemulihan_Nadi" in wb:
        order.append("01_Pemulihan_Nadi")
    order += ["02_Postur_Kerja", "03_Waktu_Kerja"]
    if "03_Perbandingan_Metode" in wb:
        order.append("03_Perbandingan_Metode")
    order += ["03_Data_Mentah",
             "04_Antropometri_32D", "04_Data_Mentah", "05_Antropometri_Lanjutan", "05_Data_Mentah",
             "06_08_Lingkungan_Kerja", "06_08_Data_Mentah", "09_MMH"]
    if "09_Simulasi" in wb:
        order.append("09_Simulasi")
    wb._sheets = [wb[name] for name in order]

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
