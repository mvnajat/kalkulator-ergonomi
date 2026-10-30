"""Validated Excel observation templates for modules 03 through 08."""

import io
import math

from openpyxl import Workbook, load_workbook


TABLES = {
    "03": ("df_siklus", "Kode", tuple(f"S{i}" for i in range(1, 11))),
    "04": ("df_32", "Kode", tuple(f"P{i}" for i in range(1, 11))),
    "05": ("df_3dim", "Dimensi", tuple(f"S{i}" for i in range(1, 11))),
}
ENVIRONMENT = {
    "06": ("dingin_suhu", "panas_suhu", 10, 50),
    "07": ("normal_db", "bising_db", 30, 120),
    "08": ("normal_lux", "gelap_lux", 1, 2000),
}
MAX_XLSX_BYTES = 2 * 1024 * 1024


def _module_number(module: str) -> str:
    module = str(module).zfill(2)
    if module not in TABLES and module not in ENVIRONMENT:
        raise ValueError("Impor Excel tersedia untuk Modul 03-08.")
    return module


def _rows_for_module(module: str, state: dict) -> list[dict]:
    # The template builder receives the whole session; the importer may receive
    # the one module only so invalid files cannot mutate the session.
    data = state.get(f"modul_{module}", state)
    if module in TABLES:
        return [dict(row) for row in data[TABLES[module][0]]]
    a, b, *_ = ENVIRONMENT[module]
    return [{"Percobaan": i + 1, a: data[a][i], b: data[b][i]} for i in range(3)]


def make_excel_template(module: str, state: dict) -> bytes:
    module = _module_number(module)
    rows = _rows_for_module(module, state)
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    if not rows:
        raise ValueError("Belum ada baris data untuk dijadikan template.")
    columns = list(rows[0])
    ws.append(columns)
    for row in rows:
        ws.append([row.get(col) for col in columns])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    from openpyxl.styles import PatternFill, Font
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="166534")
        cell.font = Font(color="FFFFFF", bold=True)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(42, max(14, max(len(str(c.value or "")) for c in col) + 2))
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def _number(value, *, maximum: float, label: str) -> float:
    if isinstance(value, bool) or value is None or str(value).strip() == "":
        raise ValueError(f"{label}: angka wajib diisi.")
    if isinstance(value, str) and value.startswith("="):
        raise ValueError(f"{label}: rumus Excel tidak diterima; gunakan nilai hasil pengukuran.")
    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{label}: bukan angka.") from exc
    if not math.isfinite(number) or number <= 0 or number > maximum:
        raise ValueError(f"{label}: gunakan angka positif sampai {maximum:g}.")
    return number


def read_excel_observations(module: str, blob: bytes, current: dict) -> dict:
    """Return a NEW module dict, without mutating the active session on errors."""
    module = _module_number(module)
    if not blob or len(blob) > MAX_XLSX_BYTES:
        raise ValueError("File Excel kosong atau lebih besar dari 2 MB.")
    try:
        wb = load_workbook(io.BytesIO(blob), read_only=True, data_only=False)
    except Exception as exc:
        raise ValueError("File bukan workbook .xlsx yang dapat dibaca.") from exc
    try:
        if "Data" not in wb.sheetnames:
            raise ValueError("Sheet 'Data' tidak ditemukan. Gunakan template aplikasi.")
        ws = wb["Data"]
        if (ws.max_row or 0) > 151 or (ws.max_column or 0) > 24:
            raise ValueError("File terlalu besar: maksimum 150 baris dan 24 kolom.")
        values = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
    if len(values) < 2:
        raise ValueError("Sheet 'Data' tidak memiliki hasil pengukuran.")
    headers = [str(x).strip() if x is not None else "" for x in values[0]]
    if len(headers) != len(set(headers)):
        raise ValueError("Kolom Excel duplikat tidak diperbolehkan.")
    expected = list(_rows_for_module(module, current)[0])
    if headers != expected:
        raise ValueError("Kolom tidak sesuai template. Kolom wajib: " + ", ".join(expected))
    rows = [dict(zip(headers, row)) for row in values[1:] if any(v is not None for v in row)]
    if not rows:
        raise ValueError("Tidak ada baris pengamatan yang terisi.")
    result = dict(current)
    if module in TABLES:
        table_key, id_col, measure_cols = TABLES[module]
        if len(rows) > 100:
            raise ValueError("Maksimum 100 baris pengamatan.")
        seen = set()
        validated = []
        for i, row in enumerate(rows, 2):
            ident = str(row.get(id_col, "") or "").strip()
            if not ident or len(ident) > 100 or ident in seen or ident.startswith("="):
                raise ValueError(f"Baris {i}: nama/kode kosong, duplikat, atau tidak valid.")
            seen.add(ident)
            checked = {}
            for col, value in row.items():
                if col in measure_cols:
                    checked[col] = _number(value, maximum=100000.0, label=f"Baris {i} {col}")
                elif col == "Kelonggaran (cm)":
                    try:
                        number = float(value)
                    except (TypeError, ValueError) as exc:
                        raise ValueError(f"Baris {i}: kelonggaran bukan angka.") from exc
                    if not math.isfinite(number) or not 0 <= number <= 100:
                        raise ValueError(f"Baris {i}: kelonggaran harus 0-100 cm.")
                    checked[col] = number
                elif col == "Persentil Acuan":
                    if value not in ("P5", "P50", "P95"):
                        raise ValueError(f"Baris {i}: persentil hanya P5, P50, atau P95.")
                    checked[col] = value
                else:
                    item = str(value or "").strip()
                    if len(item) > 500 or item.startswith("="):
                        raise ValueError(f"Baris {i} {col}: teks tidak valid.")
                    checked[col] = item
            validated.append(checked)
        result[table_key] = validated
    else:
        a, b, min_bound, max_bound = ENVIRONMENT[module]
        if len(rows) != 3:
            raise ValueError("Modul lingkungan membutuhkan tepat 3 baris percobaan.")
        first, second = [], []
        for i, row in enumerate(rows, 1):
            if row["Percobaan"] != i:
                raise ValueError("Nomor percobaan harus berurutan 1 sampai 3.")
            x = _number(row[a], maximum=max_bound, label=f"{a} percobaan {i}")
            y = _number(row[b], maximum=max_bound, label=f"{b} percobaan {i}")
            if module == "06" and x > 45:
                raise ValueError("Temperatur dingin maksimum 45 °C sesuai kontrol modul.")
            if x < min_bound or y < min_bound:
                raise ValueError(f"Modul {module}: nilai minimum {min_bound}.")
            first.append(x)
            second.append(y)
        result[a], result[b] = first, second
    return result
