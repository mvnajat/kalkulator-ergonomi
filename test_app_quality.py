import io
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

from utils.advanced_features import (
    hitung_niosh_lifting, bandingkan_skenario_niosh,
    deteksi_postur_mediapipe, simpan_ke_database_antropometri,
)
from utils.calculations import hitung_antropometri_statistik, hitung_waktu_kerja
from utils.exporter import export_session_to_excel


class AppQualityTests(unittest.TestCase):
    def niosh(self, **kw):
        kwargs = {"h": 32, "v": 60, "d": 45, "a": 15, "f": 2,
                  "c_type": "Sedang (Fair)", "berat_beban": 16, "durasi": "≤ 1 jam"}
        kwargs.update(kw)
        return hitung_niosh_lifting(**kwargs)

    def test_niosh_frequency_uses_duration_and_height(self):
        self.assertEqual(self.niosh(f=12)["fm"], 0.37)
        self.assertEqual(self.niosh(f=10)["fm"], 0.45)
        self.assertEqual(self.niosh(f=2.5)["fm"], 0.90)
        self.assertEqual(self.niosh(durasi="> 1–2 jam")["fm"], 0.84)
        self.assertEqual(self.niosh(durasi="> 2–8 jam")["fm"], 0.65)
        self.assertEqual(self.niosh(v=75)["cm"], 1.0)
        self.assertEqual(self.niosh(v=60)["cm"], 0.95)

    def test_niosh_out_of_scope_never_reports_safe_zero_li(self):
        for kw in ({"h": 70}, {"durasi": "> 2–8 jam", "f": 12}, {"a": 136}):
            with self.subTest(kw=kw):
                result = self.niosh(**kw)
                self.assertEqual(result["rwl"], 0)
                self.assertIsNone(result["li"])
        with self.assertRaises(ValueError):
            self.niosh(f=float("nan"))

    def test_niosh_compares_proposed_work_without_changing_original(self):
        original = {"h_cm": 50, "v_cm": 60, "d_cm": 45, "a_deg": 15,
                    "f_lift": 2, "coupling": "Sedang (Fair)",
                    "berat_beban": 16, "durasi_angkat": "≤ 1 jam"}
        proposed = {**original, "h_cm": 25, "v_cm": 75, "a_deg": 0,
                    "f_lift": 1, "berat_beban": 12, "coupling": "Baik (Good)"}
        snapshot = dict(original)
        result = bandingkan_skenario_niosh(original, proposed)
        self.assertEqual(original, snapshot)
        self.assertLess(result["usulan"]["li"], result["awal"]["li"])
        self.assertGreater(result["usulan"]["rwl"], result["awal"]["rwl"])
        self.assertGreater(result["perubahan_li_pct"], 0)
        self.assertTrue(result["li_usulan_le_1"])

    def test_niosh_comparison_never_fabricates_percentage_outside_scope(self):
        original = {"h_cm": 70, "v_cm": 60, "d_cm": 45, "a_deg": 15,
                    "f_lift": 2, "coupling": "Sedang (Fair)",
                    "berat_beban": 16, "durasi_angkat": "≤ 1 jam"}
        comparison = bandingkan_skenario_niosh(original, {**original, "h_cm": 25})
        self.assertIsNone(comparison["awal"]["li"])
        self.assertIsNone(comparison["perubahan_li_pct"])
        self.assertIn("luar cakupan", comparison["status"])

    def test_niosh_uses_unrounded_rwl_to_calculate_li(self):
        value = self.niosh(h=32.17, v=72.2, d=45.3, a=12.3, f=2.5)
        self.assertGreater(value["rwl"], 0)
        self.assertAlmostEqual(value["li"], round(16 / (
            23 * (25 / 32.17) * (1 - 0.003 * 2.8)
            * (0.82 + 4.5 / 45.3) * (1 - 0.0032 * 12.3)
            * 0.895 * 0.95), 3), places=3)

    def test_samples_reject_infinite_and_insufficient_n(self):
        result = hitung_antropometri_statistik([41, None, "", 0])
        self.assertIsNone(result["shapiro_p"])
        self.assertEqual(result["cukup"], "Belum Cukup")
        self.assertIn("Belum Cukup", hitung_waktu_kerja([4, None], 0.1, 15)["cukup"])
        with self.assertRaises(ValueError):
            hitung_antropometri_statistik([42, float("inf")])

    def test_time_study_rejects_invalid_allowance_precision_and_rating(self):
        for kwargs in ({"allowance_pct": 100}, {"s_level": 0},
                       {"k_level": -1}, {"p_westinghouse": -1.0},
                       {"allowance_pct": float("inf")}):
            with self.subTest(kwargs=kwargs), self.assertRaises(ValueError):
                params = {"p_westinghouse": 0.06, "allowance_pct": 15}
                params.update(kwargs)
                hitung_waktu_kerja([4, 5, 6], **params)

    def test_blank_image_does_not_invent_human_posture(self):
        from PIL import Image

        pic = io.BytesIO()
        Image.new("RGB", (40, 40), "white").save(pic, format="PNG")
        pic.seek(0)
        with patch("utils.advanced_features.HAS_MEDIAPIPE", False):
            result = deteksi_postur_mediapipe(pic)
        self.assertFalse(result["valid_detection"])
        self.assertIsNone(result["skor_rekomendasi"])
        self.assertEqual(result["sudut"], {})

    def test_excel_covers_modul09_and_raw_samples_with_formulas(self):
        state = {
            "identitas": {"kelas": "J", "kelompok": "Kelompok 1"},
            "modul_09": {"f_lift": 12, "durasi_angkat": "> 2–8 jam", "coupling": "Sedang (Fair)"},
            "modul_03": {"df_siklus": [{"Kode": "EL01", "Elemen Kerja": "Merakit", "S1": 4.0, "S2": 5.0}]},
            "modul_04": {"df_32": [{"Kode": "D01", "Dimensi Tubuh": "Tinggi", "P1": 160.0, "P2": 170.0}]},
            "modul_05": {"df_3dim": [{"Dimensi": "TPD", "S1": 40.0, "S2": 42.0, "Persentil Acuan": "P5"}]},
        }
        results = {"m9_summary": {"rwl": 0, "li": None, "kategori": "Di luar batas", "f_kompresi_n": 2500}}
        workbook = load_workbook(io.BytesIO(export_session_to_excel(state, results)))
        self.assertEqual(len(workbook.sheetnames), 12)
        for tab in ("04_Data_Mentah", "03_Data_Mentah", "05_Data_Mentah", "06_08_Data_Mentah", "09_MMH"):
            self.assertIn(tab, workbook.sheetnames)
        self.assertEqual(workbook["04_Data_Mentah"]["C5"].value, 160.0)
        self.assertTrue(workbook["04_Data_Mentah"]["M5"].value.startswith("=IF("))
        self.assertEqual(workbook["09_MMH"]["B10"].value, "> 2–8 jam")
        self.assertIn("NIOSH Lifting Index", [workbook["09_MMH"].cell(row, 1).value for row in range(5, 24)])

        state["identitas"]["kelas"] = '=HYPERLINK("https://example.com", "click")'
        workbook = load_workbook(io.BytesIO(export_session_to_excel(state, results)))
        self.assertEqual(workbook["00_Identitas"]["B5"].data_type, "s")

    def test_excel_includes_niosh_scenario_when_available(self):
        original = {"h_cm": 50, "v_cm": 60, "d_cm": 45, "a_deg": 15,
                    "f_lift": 2, "coupling": "Sedang (Fair)",
                    "berat_beban": 16, "durasi_angkat": "≤ 1 jam"}
        proposal = {**original, "h_cm": 25, "berat_beban": 12}
        comparison = bandingkan_skenario_niosh(original, proposal)
        workbook = load_workbook(io.BytesIO(export_session_to_excel(
            {"modul_09": original, "modul_09_simulasi": proposal},
            {"m9_simulasi_summary": comparison},
        )))
        self.assertIn("09_Simulasi", workbook.sheetnames)
        self.assertEqual(workbook["09_Simulasi"]["B5"].value, 50)
        self.assertEqual(workbook["09_Simulasi"]["C5"].value, 25)

    def test_database_refuses_empty_or_invalid_records(self):
        with self.assertRaises(ValueError):
            simpan_ke_database_antropometri("", "K1", [])
        with self.assertRaises(ValueError):
            simpan_ke_database_antropometri("J", "K1", [])
        with self.assertRaises(ValueError):
            simpan_ke_database_antropometri("J", "K1", [{"Kode": "D01", "Dimensi Tubuh": "Tinggi", "P1": float("inf")}])


if __name__ == "__main__":
    unittest.main()
